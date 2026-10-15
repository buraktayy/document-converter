"""PDF kaynaklı/hedefli dönüştürmeler — Office gerekmez, saf Python.

PDF -> Word burada pdf2docx ile yapılır (bkz. office_com.py'deki not: Word'ün
yerleşik PDF-aç-dönüştür özelliği otomasyonla bastırılamayan bir onay
diyaloğu yüzünden COM üzerinden donuyor, bu yüzden tercih edilmedi).
"""
from __future__ import annotations

import io
from pathlib import Path

import fitz  # PyMuPDF
import openpyxl
import pdfplumber
from PIL import Image
from pdf2docx import Converter as Pdf2DocxConverter
from pptx import Presentation

from .. import fileops, security
from .types import ConversionOptions, ConversionResult

SLIDE_WIDTH_EMU_169 = 12192000  # 13.333in genişlik, 16:9 baz tasarım


def pdf_to_word(source: Path, dest: Path, options: ConversionOptions) -> ConversionResult:
    warning = ""
    if not security.pdf_has_text_layer(source):
        warning = " (Uyarı: Bu PDF taranmış görünüyor, metin az veya hiç çıkmamış olabilir.)"
    try:
        with fileops.atomic_write_target(dest) as tmp_path:
            converter = Pdf2DocxConverter(str(source))
            try:
                converter.convert(str(tmp_path))
            finally:
                converter.close()
        return ConversionResult(
            success=True, output_paths=[dest], message=f"{dest.name} oluşturuldu.{warning}"
        )
    except Exception as exc:  # noqa: BLE001
        return ConversionResult(
            success=False,
            message="PDF Word'e dönüştürülemedi (dosya bozuk veya çok karmaşık bir yapıya sahip olabilir).",
            error_detail=repr(exc),
        )


def _save_pixmap(pix: "fitz.Pixmap", dest_path: Path, image_ext: str, jpeg_quality: int = 92) -> None:
    if image_ext == "png":
        pix.save(str(dest_path))
        return
    png_bytes = pix.tobytes("png")
    with Image.open(io.BytesIO(png_bytes)) as img:
        img.convert("RGB").save(str(dest_path), format="JPEG", quality=jpeg_quality)


def pdf_to_images(source: Path, dest: Path, options: ConversionOptions) -> ConversionResult:
    """dest, hedef klasörü (dest.parent) ve istenen formatı (dest.suffix) taşıyan
    bir 'şablon' yoldur — birden fazla sayfa çıktısı üretildiği için dest'in
    kendisi doğrudan kullanılmaz, her sayfa için ayrı bir dosya adı üretilir."""
    image_ext = security.normalize_ext(dest.suffix)
    output_dir = dest.parent
    try:
        output_paths: list[Path] = []
        with fitz.open(source) as doc:
            if doc.page_count == 0:
                return ConversionResult(success=False, message="PDF boş görünüyor.")
            zoom = options.dpi / 72
            matrix = fitz.Matrix(zoom, zoom)
            multi_page = doc.page_count > 1
            for i, page in enumerate(doc, start=1):
                pix = page.get_pixmap(matrix=matrix)
                suffix = f"_s{i}" if multi_page else ""
                dest = fileops.build_output_path(source, output_dir, image_ext, suffix=suffix)
                with fileops.atomic_write_target(dest) as tmp_path:
                    _save_pixmap(pix, tmp_path, image_ext)
                output_paths.append(dest)
        return ConversionResult(
            success=True,
            output_paths=output_paths,
            message=f"{len(output_paths)} sayfa görsele dönüştürüldü.",
        )
    except Exception as exc:  # noqa: BLE001
        return ConversionResult(success=False, message="PDF görsellere dönüştürülemedi.", error_detail=repr(exc))


def pdf_to_excel(source: Path, dest: Path, options: ConversionOptions) -> ConversionResult:
    try:
        wb = openpyxl.Workbook()
        wb.remove(wb.active)
        any_content = False
        with pdfplumber.open(source) as pdf:
            if len(pdf.pages) == 0:
                return ConversionResult(success=False, message="PDF boş görünüyor.")
            for i, page in enumerate(pdf.pages, start=1):
                ws = wb.create_sheet(title=f"Sayfa {i}"[:31])
                row_cursor = 1
                tables = page.extract_tables()
                if tables:
                    for table in tables:
                        for row in table:
                            for col_idx, value in enumerate(row, start=1):
                                ws.cell(row=row_cursor, column=col_idx, value=value)
                            row_cursor += 1
                        row_cursor += 1
                        any_content = True
                else:
                    text = page.extract_text() or ""
                    for line in text.splitlines():
                        ws.cell(row=row_cursor, column=1, value=line)
                        row_cursor += 1
                        if line.strip():
                            any_content = True
        if not any_content:
            return ConversionResult(success=False, message="PDF içinde tablo veya metin bulunamadı.")
        with fileops.atomic_write_target(dest) as tmp_path:
            wb.save(str(tmp_path))
        return ConversionResult(success=True, output_paths=[dest], message=f"{dest.name} oluşturuldu.")
    except Exception as exc:  # noqa: BLE001
        return ConversionResult(success=False, message="PDF Excel'e dönüştürülemedi.", error_detail=repr(exc))


def pdf_to_powerpoint(source: Path, dest: Path, options: ConversionOptions) -> ConversionResult:
    try:
        with fitz.open(source) as doc:
            if doc.page_count == 0:
                return ConversionResult(success=False, message="PDF boş görünüyor.")

            zoom = options.dpi / 72
            matrix = fitz.Matrix(zoom, zoom)

            first_rect = doc[0].rect
            deck_ratio = first_rect.width / first_rect.height
            slide_width = SLIDE_WIDTH_EMU_169
            slide_height = int(slide_width / deck_ratio)

            prs = Presentation()
            prs.slide_width = slide_width
            prs.slide_height = slide_height
            blank_layout = prs.slide_layouts[6]

            for page in doc:
                pix = page.get_pixmap(matrix=matrix)
                img_bytes = pix.tobytes("png")
                page_ratio = page.rect.width / page.rect.height

                if abs(page_ratio - deck_ratio) < 0.01:
                    left, top, width, height = 0, 0, slide_width, slide_height
                elif page_ratio > deck_ratio:
                    width, height = slide_width, int(slide_width / page_ratio)
                    left, top = 0, (slide_height - height) // 2
                else:
                    height, width = slide_height, int(slide_height * page_ratio)
                    top, left = 0, (slide_width - width) // 2

                slide = prs.slides.add_slide(blank_layout)
                slide.shapes.add_picture(io.BytesIO(img_bytes), left, top, width=width, height=height)

        with fileops.atomic_write_target(dest) as tmp_path:
            prs.save(str(tmp_path))
        return ConversionResult(success=True, output_paths=[dest], message=f"{dest.name} oluşturuldu.")
    except Exception as exc:  # noqa: BLE001
        return ConversionResult(
            success=False, message="PDF PowerPoint'e dönüştürülemedi.", error_detail=repr(exc)
        )
